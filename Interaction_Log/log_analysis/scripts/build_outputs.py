#!/usr/bin/env python3
"""Turn the coded interaction logs into the Excel workbook and the theme chart.

Inputs (in ../coding_data):
    logs_parsed.csv       one row per ';'-separated log      (parse_log.py)
    primary_codes.txt     one category per log               (hand-coded)
    codebook.csv          category -> theme

Outputs (in the interaction_log folder):
    Interaction_Log_Thematic_Analysis.xlsx
    charts/themes_count.png  (bars labelled "count (% of group)")

    python build_outputs.py

Needs: pandas, openpyxl, matplotlib, seaborn.
"""
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
import seaborn as sns
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "coding_data"
CHARTS = ROOT / "charts"
CHARTS.mkdir(exist_ok=True)
OUT_XLSX = ROOT / "Interaction_Log_Thematic_Analysis.xlsx"

GROUPS = {"E": "Experiment (E)", "C": "Control (C)"}
COLORS = {"Control (C)": "#4C72B0", "Experiment (E)": "#DD8452"}  # seaborn "deep" blue/orange


def load():
    cb = pd.read_csv(DATA / "codebook.csv")
    logs = pd.read_csv(DATA / "logs_parsed.csv")
    codes = pd.read_csv(DATA / "primary_codes.txt", sep=" ", comment="#", names=["log_id", "code"])
    df = logs.merge(codes, on="log_id", how="left")
    missing = df.code.isna().sum()
    unknown = set(df.code.dropna()) - set(cb.code)
    assert not missing and not unknown, f"{missing} uncoded logs, unknown codes {unknown}"
    return cb, df.merge(cb, on="code").rename(columns={"log": "text"})


def theme_table(df, cb):
    themes = cb.drop_duplicates("theme_id")[["theme_id", "theme"]]
    n = df.groupby(["theme_id", "group"]).size().unstack(fill_value=0)
    n = n.reindex(themes.theme_id, fill_value=0).reindex(columns=["E", "C"], fill_value=0)
    ppl = df.groupby(["theme_id", "group"]).pid.nunique().unstack(fill_value=0)
    ppl = ppl.reindex(themes.theme_id, fill_value=0).reindex(columns=["E", "C"], fill_value=0)
    out = themes.set_index("theme_id")
    out["tasks"] = df.groupby("theme_id").task.unique().reindex(out.index).map(task_note)
    for g in ("E", "C"):
        out[f"{g} count"] = n[g]
        out[f"{g} % of group"] = (100 * n[g] / n[g].sum()).round(1)
        out[f"{g} participants (of 15)"] = ppl[g]
        out[f"{g} mean per participant"] = (n[g] / 15).round(2)
    return out.reset_index()


def category_table(df, cb):
    n = df.groupby(["code", "group"]).size().unstack(fill_value=0).reindex(columns=["E", "C"], fill_value=0)
    out = cb.set_index("code").join(n).fillna(0)
    for g in ("E", "C"):
        out[g] = out[g].astype(int)
        out[f"{g} %"] = (100 * out[g] / out[g].sum()).round(1)
    return out.reset_index()[["theme_id", "theme", "code", "category", "E", "E %", "C", "C %"]]


def category_lists(df, cb, group):
    """One row per category: 'log (P#, T#); log (P#, T#); ...' — the layout the analysis asked for."""
    sub = df[df.group == group].sort_values(["pnum", "task", "log_id"])
    rows = []
    for r in cb.itertuples():
        items = sub[sub.code == r.code]
        rows.append({"Theme": r.theme, "Code": r.code, "Category": r.category, "n": len(items),
                     "Participants": ", ".join(dict.fromkeys(items.pid)),
                     "Behavioral logs": "; ".join(f"{t} ({p}, T{k})" for t, p, k in
                                                  zip(items.text, items.pid, items.task))})
    return pd.DataFrame(rows)


# Two-line tick labels for the chart (themes not listed stay on one line).
TICK_LABELS = {
    "Read Task Instructions": "Read Task\nInstructions",
    "Review Task Artifacts": "Review Task\nArtifacts",
    "Set Up Workspace": "Set Up\nWorkspace",
    "Explore Application": "Explore\nApplication",
    "Identify Problems Manually": "Identify Problems\nManually",
    "Prompt AI via Text Chat": "Prompt AI\nvia Text Chat",
    "Interact via Real Time": "Interact\nvia Real Time",
    "Review AI Response": "Review AI\nResponse",
    "Write & Integrate Code": "Write &\nIntegrate Code",
    "Test & Verify Implementation": "Test &\nVerify Implementation",
    "Write Problem Report": "Write Problem\nReport",
}


def task_note(tasks):
    """'(Task 3)' / '(Tasks 1 & 2)' for a theme seen only in some tasks; '' if it spans all tasks."""
    tasks = sorted(tasks)
    if tasks == [1, 2, 3]:
        return ""
    return f"(Task {tasks[0]})" if len(tasks) == 1 else f"(Tasks {' & '.join(map(str, tasks))})"


def centered_note(fig, bold, rest, y, fontsize):
    """One line of figure text, horizontally centered, whose first words are bold."""
    renderer = fig.canvas.get_renderer()
    head = fig.text(0, y, bold, fontsize=fontsize, fontweight="bold", va="bottom")
    tail = fig.text(0, y, rest, fontsize=fontsize, va="bottom")
    to_fig = fig.transFigure.inverted()
    w_head = to_fig.transform((head.get_window_extent(renderer).width, 0))[0] - to_fig.transform((0, 0))[0]
    w_tail = to_fig.transform((tail.get_window_extent(renderer).width, 0))[0] - to_fig.transform((0, 0))[0]
    x0 = 0.5 - (w_head + w_tail) / 2
    head.set_x(x0)
    tail.set_x(x0 + w_head)


def plot(tt):
    """Bars show counts; each bar is labelled 'count (% of that group's behaviors)'."""
    long = pd.concat([
        pd.DataFrame({"Theme": tt.theme, "Group": GROUPS[g], "value": tt[f"{g} count"]})
        for g in ("C", "E")])
    sns.set_theme(style="whitegrid", context="talk")
    fig, ax = plt.subplots(figsize=(22, 7.5))
    sns.barplot(data=long, x="Theme", y="value", hue="Group", palette=COLORS,
                hue_order=list(COLORS), ax=ax, saturation=1)
    for cont, g in zip(ax.containers, ("C", "E")):
        labels = [f"{n}\n({p:.1f}%)" for n, p in zip(tt[f"{g} count"], tt[f"{g} % of group"])]
        ax.bar_label(cont, labels=labels, fontsize=11, padding=2, linespacing=1.0)
    ax.set_ylim(0, long.value.max() * 1.18)
    ax.set_xticks(range(len(tt)))
    labels = [TICK_LABELS.get(t, t) + (f"\n{note}" if note else "") for t, note in zip(tt.theme, tt.tasks)]
    ax.set_xticklabels(labels, fontsize=13)
    ax.set_xlabel("Theme")
    ax.set_ylabel("Number of behavioral logs")
    totals = f"E = {int(tt['E count'].sum())}, C = {int(tt['C count'].sum())} logs"
    ax.set_title(f"Number (and % of group) of Behaviors per Theme ({totals})", fontsize=17)
    ax.legend(title="Group", loc="upper right", fontsize=13, title_fontsize=13)
    fig.tight_layout(rect=(0, 0.05, 1, 1))  # leave room for the note below the axis label
    centered_note(fig, "Misc.", " represents activities including interacting with the experimenter, "
                  "reflecting on or planning AI usage (think-aloud), and timing.", y=0.015, fontsize=13)
    path = CHARTS / "themes_count.png"
    fig.savefig(path, dpi=200)
    plt.close(fig)
    return path


def style_sheet(ws, widths, wrap_cols=()):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE5F0")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=cell.column in wrap_cols, vertical="top")
    ws.freeze_panes = "A2"


def main():
    cb, df = load()
    tt = theme_table(df, cb)
    notes = pd.DataFrame({"Item": [
        "Source", "Groups", "Unit", "Coding", "Themes", "Counts", "P17 update", "Tasks", "Scripts"], "Detail": [
        "Interaction Log.pdf, updated version with P17 completed (observer notes from the 30-participant user study).",
        "Experiment (E) = odd participants P1, P3 … P29; Control (C) = even participants P2, P4 … P30.",
        f"Each ';'-separated note is one behavioral log ({len(df)} logs) and gets exactly one category: "
        "the log's main action. A ';' inside parentheses (an example list) does not start a new log. "
        "P9's notes put one action per line instead of using ';', so each line is one log.",
        "Inductive open coding into 53 categories, grouped into 12 themes (see Codebook). "
        "Codes were assigned by reading every log.",
        "; ".join(cb.drop_duplicates("theme_id").theme),
        "'% of group' = theme count / all behaviors of that group, so the two groups are comparable even "
        "though note length differs by participant.",
        "P17's notes were completed in the updated PDF (2026-10-21); all earlier P17 logs were replaced with the "
        "new ones and coded fresh. All other participants' logs are unchanged.",
        "Task number and the AI feature noted for each task are kept in the 'All logs' sheet.",
        "scripts/parse_log.py -> scripts/build_outputs.py (codes in coding_data/primary_codes.txt)."]})

    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as xw:
        notes.to_excel(xw, sheet_name="README", index=False)
        style_sheet(xw.sheets["README"], [18, 140], wrap_cols=(2,))

        n = df.groupby(["code", "group"]).size().unstack(fill_value=0)
        cbx = cb.copy()
        for g in ("E", "C"):
            cbx[g] = cbx.code.map(n.get(g, pd.Series(dtype=int))).fillna(0).astype(int)
        cbx.to_excel(xw, sheet_name="Codebook", index=False)
        style_sheet(xw.sheets["Codebook"], [7, 58, 9, 32, 8, 8], wrap_cols=(2,))

        tt.to_excel(xw, sheet_name="Theme counts", index=False)
        style_sheet(xw.sheets["Theme counts"], [9, 30, 13] + [12] * 8)

        category_table(df, cb).to_excel(xw, sheet_name="Category counts", index=False)
        style_sheet(xw.sheets["Category counts"], [9, 30, 7, 58, 8, 8, 8, 8])

        for g in ("E", "C"):
            sheet = f"{g} categories"
            category_lists(df, cb, g).to_excel(xw, sheet_name=sheet, index=False)
            style_sheet(xw.sheets[sheet], [26, 7, 40, 6, 30, 160], wrap_cols=(3, 5, 6))

        cols = ["log_id", "pid", "group", "task", "condition", "text", "code", "category", "theme"]
        df.sort_values(["pnum", "task", "seq"])[cols].to_excel(xw, sheet_name="All logs", index=False)
        style_sheet(xw.sheets["All logs"], [7, 6, 6, 5, 30, 90, 6, 40, 28], wrap_cols=(6,))

    print(OUT_XLSX)
    print(plot(tt))
    print(tt.drop(columns="theme_id").to_string(index=False))


if __name__ == "__main__":
    main()
